#!/usr/bin/env python3
"""Memoria editorial única: HTML offline y PDF desde los mismos bloques.

Requiere numpy, scipy, openpyxl (solo lectura), matplotlib, reportlab y
latex2mathml. Los adjuntos originales permanecen fuera del control de versiones.
"""
from pathlib import Path
from io import BytesIO
from html import escape
import base64
import json
import hashlib
import sys
import numpy as np
from scipy.optimize import brentq, root
from scipy.stats import norm, chi2
from openpyxl import load_workbook
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.mathtext import math_to_image
from matplotlib.font_manager import FontProperties
from matplotlib.ticker import PercentFormatter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, Image, KeepTogether
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image as PILImage
import latex2mathml.converter

ROOT = Path(__file__).resolve().parents[1]
PROC = ROOT / 'procesos-estocasticos'
OUT = ROOT / 'output'
TMP = ROOT / 'tmp' / 'output_procesos'
OUT.mkdir(exist_ok=True)
TMP.mkdir(parents=True, exist_ok=True)
RATINGS = ['Aaa','Aa','A','Baa','Ba','B','Caa-C']
STATES = RATINGS + ['Default']
source = PROC / '0-enunciado' / 'matriz-ratings.xlsx'
ws = load_workbook(source, data_only=True).active
Q = np.array([[ws.cell(i,j).value for j in range(3,11)] for i in range(3,12)],float)
P = Q[[0,1,2,3,4,5,6,8]]
assert P.shape == (8,8) and np.all(P >= 0)
assert np.max(abs(P.sum(1)-1)) < 1e-14
assert np.array_equal(P[-1], [0,0,0,0,0,0,0,1])
F = np.array([np.linalg.matrix_power(P,n)[:7,-1] for n in range(1,26)])
f = np.diff(np.vstack([np.zeros(7),F]),axis=0)
counts = np.array([136,694,1298,1175,578,817,304])
w = counts / counts.sum()
g = f @ w
P25 = np.linalg.matrix_power(P,25)
assert counts.sum()==5002 and np.all(f >= 0)
assert np.allclose(f.sum(0),F[-1])
bookpath = PROC / '2-ito-martingalas/resultados/ejercicio2_simulacion_Mt.xlsx'
book = load_workbook(bookpath, data_only=True)
M = np.array([book['Replica_t1'].cell(i,4).value for i in range(3,5003)])
Z = np.array([book['Replica_t1'].cell(i,2).value for i in range(3,5003)])
assert np.allclose(M,(Z**2-1)/3,atol=1e-14)
mean, var = M.mean(), M.var(ddof=1)
se = np.sqrt(var/len(M))
summary = {row[3]: row[5] for row in book['Resumen'].values if len(row) >= 6 and row[3]}
assert abs(mean-summary['Media de M₁']) < 1e-14
assert abs(var-summary['Varianza muestral']) < 1e-14
times = np.array([1.,1.25,2.])
prices = np.array([9.55,12.48,23.36])
def call(T,V):
    d1=(.01*T+.5*V)/np.sqrt(V)
    return 100*norm.cdf(d1)-100*np.exp(-.01*T)*norm.cdf(d1-np.sqrt(V))
Vs=np.array([brentq(lambda v:call(t,v)-c,1e-12,2,xtol=1e-14) for t,c in zip(times,prices)])
def V(t,abc):
    a,b,c=abc
    return a*a*t**5/5+a*b*t**4/2+(b*b+2*a*c)*t**3/3+b*c*t*t+c*c*t
abc=root(lambda x:V(times,x)-Vs,[.05,.2,.1],tol=1e-11).x
a,b,c=abc
assert np.max(abs(V(times,abc)-Vs)) < 1e-12
assert min(a,b,c)>0
fit=np.array([call(t,V(t,abc)) for t in times])
assert np.max(abs(fit-prices))<1e-10

def pct(x,d=4): return f'{100*x:.{d}f}'.replace('.',',')+' %'
def num(x,d=4): return f'{x:.{d}f}'.replace('.',',')
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,'axes.edgecolor':'#bdc7c7','axes.labelcolor':'#334644','text.color':'#1e302e','xtick.color':'#526360','ytick.color':'#526360','grid.alpha':.18})
fig,axes=plt.subplots(3,1,figsize=(6.7,6.3),layout='constrained')
for ax,idx,title in zip(axes,[[0,1,2],[3,4],[5,6]],['Ratings altos','Ratings intermedios','Ratings bajos']):
    for i,col in zip(idx,['#175b54','#8b654b','#769793']):ax.plot(range(1,26),f[:,i],label=RATINGS[i],color=col,lw=1.8)
    ax.set_title(title,loc='left',fontsize=11)
    ax.yaxis.set_major_formatter(PercentFormatter(1))
    ax.set_xlim(1,25); ax.grid(axis='y'); ax.legend(frameon=False,loc='upper right',ncol=3)
axes[-1].set_xlabel('Año del primer default')
fig.savefig(TMP/'default.png',dpi=190);plt.close(fig)
fig,ax=plt.subplots(figsize=(6.7,2.9),layout='constrained')
xx=np.linspace(-1/3,3.7,700)
ax.plot(xx,chi2.cdf(3*xx+1,1),color='#175b54',lw=2,label='Distribución teórica')
ordered=np.sort(M)
ax.plot(ordered,np.arange(1,len(M)+1)/len(M),color='#8b654b',ls='--',lw=1.4,label='5.000 réplicas del Excel')
ax.set(xlabel='Valor de M₁',ylabel='Probabilidad acumulada',ylim=(0,1.02),xlim=(-.4,3.7));ax.legend(frameon=False);ax.grid(axis='y')
fig.savefig(TMP/'cdf.png',dpi=190);plt.close(fig)
fig,ax=plt.subplots(figsize=(6.7,2.8),layout='constrained')
tt=np.linspace(0,2,250)
ax.plot(tt,a*tt**2+b*tt+c,color='#175b54',lw=2)
ax.set(xlabel='Tiempo (años)',ylabel='Volatilidad instantánea',xlim=(0,2))
ax.yaxis.set_major_formatter(PercentFormatter(1));ax.grid(axis='y')
fig.savefig(TMP/'vol.png',dpi=190);plt.close(fig)

# Un mismo contenido alimenta las dos presentaciones. Cada página PDF es
# una unidad de lectura; en HTML se mantiene el flujo vertical continuo.
pages=[]
def page(title,kicker):
    blocks=[];pages.append((title,kicker,blocks));return blocks
def p(b,s):b.append(('p',s))
def h(b,s):b.append(('h',s))
def eq(b,tex,label):b.append(('eq',tex,label))
def note(b,s):b.append(('note',s))
def table(b,headers,rows):b.append(('table',headers,rows))
def fig(b,name,caption):b.append(('fig',name,caption))

b0=page('Procesos estocásticos','MEFC 2026 · Memoria de resolución')
p(b0,'Tres ejercicios, un mismo recorrido: definir el modelo, desarrollar el cálculo y comprobar qué significa el resultado. Esta memoria sigue el orden del enunciado y reúne las respuestas teóricas y numéricas del trabajo.')
table(b0,['Ejercicio','Pregunta central','Resultado'],[
 ['1 · Ratings','¿Cuándo llega el default?','54,1342 % acumulado a 25 años en la cohorte.'],
 ['2 · Itô','¿Cómo se compensa la deriva?','μ = σρ − 1/2; E[Mₜ] = 0.'],
 ['3 · Opciones','¿Cómo valorar con σ variable?','Call asiática ≈ 13,62 unidades monetarias.']])
h(b0,'Convenciones de lectura')
p(b0,'El tiempo se expresa en años. Las probabilidades son adimensionales; las tablas las presentan en porcentaje. E denota esperanza, Var varianza y Φ la función de distribución de una normal estándar. Un intervalo de confianza de simulación describe error numérico, no incertidumbre del modelo.')
h(b0,'Decisión sobre el dato de ratings')
p(b0,'El Excel tiene nueve filas de origen y ocho columnas de destino. El enunciado considera ocho estados: Aaa, Aa, A, Baa, Ba, B, Caa-C y Default. La profesora autoriza eliminar una de las dos filas de salida de la categoría inferior, sin inventar pesos de agregación.')
note(b0,'Se conserva la fila Caa, se elimina Ca-C y se asigna la fila conservada al estado Caa-C. Esta convención se aplica a todos los cálculos del ejercicio 1. No se añaden columnas ni se promedian filas.')
p(b0,'Fuentes: MEFC_2026_examen_procesos.pdf, matriz-ratings.xlsx y la aclaración de Maite facilitada al grupo. Los notebooks del repositorio contienen el código y las simulaciones; el Excel existente documenta la parte 2.c. Los anexos de esta memoria contienen las tablas completas de los años 1 a 25.')

b0=page('Default acumulado y primer default','Ejercicio 1 · Apartado a')
p(b0,'La matriz anual P indica cómo puede cambiar el rating en un año. Su elemento pᵢⱼ es la probabilidad de pasar del estado i al j. Usamos una cadena de Markov homogénea: la transición depende del rating actual y la misma matriz se mantiene cada año.')
p(b0,'Default es absorbente: una compañía que ha llegado a él permanece allí. Por eso, estar en default en el año n equivale a haber hecho default en algún momento hasta ese año.')
eq(b0,r'F_i(n)=(P^n)_{iD},\qquad n=1,\ldots,25','Probabilidad acumulada de default desde el rating i en el año n.')
p(b0,'D identifica Default; i identifica el rating inicial. La potencia Pⁿ encadena n transiciones anuales. Para un rating inicialmente solvente fijamos Fᵢ(0) = 0.')
p(b0,'Para obtener el primer default en un año concreto restamos el acumulado del año anterior:')
eq(b0,r'f_i(n)=F_i(n)-F_i(n-1)','Probabilidad de primer default exactamente en el año n.')
table(b0,['Rating inicial','Default en 1 año','Acumulado a 25 años'],[[r,pct(F[0,i]) if F[0,i]>=1e-6 else f'{100*F[0,i]:.3e} %',pct(F[-1,i])]for i,r in enumerate(RATINGS)])
note(b0,'Ejemplo: desde Caa-C, F(1) = 35,5200 %. El segundo año se calcula como F(2) − F(1), no como F(2) aislado. Las probabilidades f son incondicionales respecto al grupo inicial: no se dividen por los supervivientes.')
p(b0,'Comprobación: todas las fᵢ(n) son no negativas y su suma de los años 1 a 25 coincide con Fᵢ(25). Si se parte ya de Default, el acumulado es 100 % y el primer default ocurrió en el instante inicial, no en los años futuros.')

b0=page('Por qué las curvas son distintas','Ejercicio 1 · Interpretación del apartado a')
fig(b0,'default.png','Probabilidad de primer default por año y rating inicial. Cada panel tiene su propia escala vertical para permitir la lectura de cada grupo.')
p(b0,'En ratings altos, la probabilidad inicial es muy pequeña: suele ser necesario deteriorarse antes de llegar a Default. En ratings intermedios aparece una joroba; en los bajos pesa más el default temprano.')
p(b0,'Los máximos dentro de la ventana de 25 años se observan en los años 25, 24, 11, 8, 5, 2 y 1 para Aaa, Aa, A, Baa, Ba, B y Caa-C, respectivamente. En particular, B alcanza el máximo en el año 2, mientras que Caa-C lo hace en el primero.')
note(b0,'El máximo de Aaa en el año 25 es solo el mayor valor dentro de la ventana solicitada. No demuestra que sea el máximo de toda su distribución temporal.')

b0=page('La cohorte de 5.002 compañías','Ejercicio 1 · Apartado b')
p(b0,'El enunciado da 136 compañías Aaa, 694 Aa, 1.298 A, 1.175 Baa, 578 Ba, 817 B y 304 Caa-C. No ponderamos los ratings por igual: cada grupo pesa según su tamaño inicial.')
eq(b0,r'w_i=\frac{N_i}{5002},\qquad g(n)=\sum_{i=1}^{7}w_i f_i(n)','Proporción esperada de compañías que hacen default exactamente en el año n.')
p(b0,'Nᵢ es el número inicial de compañías del rating i; wᵢ su proporción y g(n) la proporción de toda la cohorte que hace default en el año n. El número esperado es 5.002 × g(n); puede no ser entero porque es una esperanza.')
h(b0,'Cálculo del primer año')
p(b0,'Se multiplican los siete tamaños iniciales por la columna Default de P y se suman. El total esperado es 113,12 compañías; al dividir por 5.002 se obtiene el 2,2614 %. La tabla siguiente resume los resultados, y el anexo C recoge los 25 años.')
table(b0,['Año','Proporción en ese año','Compañías esperadas'],[[str(n),pct(g[n-1]),num(5002*g[n-1],2)]for n in [1,2,3,5,10,15,20,25]])
note(b0,f'El máximo anual de la cohorte ocurre en el año {np.argmax(g)+1}. Sumando las proporciones de los 25 años se obtiene {pct(g.sum())}, equivalente a {num(5002*g.sum(),2)} defaults esperados.')
p(b0,'La linealidad de la esperanza permite sumar los defaults esperados aunque las compañías no fueran independientes. La independencia sería una hipótesis adicional para estudiar la dispersión del total, que aquí no se pide.')

b0=page('Distribución del rating a 25 años','Ejercicio 1 · Apartado c')
p(b0,'Si conocemos el rating inicial, utilizamos su fila de P²⁵. Si elegimos una compañía al azar de toda la cohorte y no condicionamos por su rating inicial, ponderamos esas filas con los pesos anteriores.')
eq(b0,r'\pi_{25}=\pi_0P^{25},\qquad \pi_0=(w_1,\ldots,w_7,0)','Distribución a 25 años de una compañía elegida al azar de la cohorte.')
p(b0,'Los vectores π son vectores fila de ocho probabilidades en el orden Aaa, Aa, A, Baa, Ba, B, Caa-C, Default. La tabla muestra cada distribución condicional; todas sus filas suman 100 % antes del redondeo.')
table(b0,['Origen']+STATES,[[r]+[num(x*100,3) for x in P25[i]] for i,r in enumerate(RATINGS)])
p(b0,'Todos los valores de la tabla anterior están en %. La distribución no condicionada de la cohorte es:')
table(b0,['Estado a 25 años','Probabilidad'],[[r,pct(x)] for r,x in zip(STATES,np.r_[w,0]@P25)])
h(b0,'Comprobación por simulación')
p(b0,'El notebook simula 300.000 cadenas desde cada rating, avanzando 25 veces con la fila de transición correspondiente al estado actual. Las frecuencias terminales se comparan con P²⁵: el error absoluto máximo es 0,001486, es decir, 0,1486 puntos porcentuales.')
p(b0,'La discrepancia máxima estandarizada es 2,41 errores estándar. Para una celda con probabilidad p, el error estándar de su frecuencia es √[p(1−p)/300.000]. La concordancia es compatible con variabilidad Monte Carlo; no se exige igualdad exacta entre frecuencias y probabilidades.')

b0=page('El cociente de dos procesos','Ejercicio 2 · Apartado a')
p(b0,'Buscamos la deriva μ que hace martingala a Yₜ = Xₜ/Nₜ. La idea es compensar no solo las derivas originales, sino también las correcciones de Itô y la correlación de los dos brownianos.')
eq(b0,r'dX_t=\mu X_t\,dt+\sigma X_t\,dW_t^{(1)}','Dinámica de X: deriva mu y volatilidad sigma.')
eq(b0,r'dN_t=\frac12 N_t\,dt+N_t\,dW_t^{(2)}','Dinámica de N: deriva un medio y volatilidad uno.')
p(b0,'σ y ρ son constantes dadas, con −1 ≤ ρ ≤ 1 y d⟨W⁽¹⁾,W⁽²⁾⟩ₜ = ρ dt. Suponemos valores iniciales deterministas, X₀ ≠ 0 y N₀ > 0; así Nₜ no se anula y el cociente está definido. Si X₀ = 0, el cociente es siempre cero y cualquier μ sirve.')
h(b0,'Aplicación de Itô')
p(b0,'Para la función h(x,n) = x/n, las derivadas que intervienen son hₓ = 1/n, hₙ = −x/n², hₓₓ = 0, hₙₙ = 2x/n³ y hₓₙ = −1/n². El término cruzado usa d⟨X,N⟩ₜ = σρXₜNₜ dt.')
eq(b0,r'dY_t=Y_t\left[(\mu+\frac12-\sigma\rho)dt+\sigma dW_t^{(1)}-dW_t^{(2)}\right]','Dinámica del cociente, incluida la corrección de covariación.')
p(b0,'La deriva total es μ − 1/2 + 1 − σρ. El +1 procede de la segunda derivada respecto a N; el −σρ, de la covariación. Al igualarla a cero:')
eq(b0,r'\boxed{\mu=\sigma\rho-\frac12}','Solución: mu igual a sigma por rho menos un medio.')
h(b0,'Por qué es una martingala verdadera')
p(b0,'Sea v = σ² + 1 − 2σρ ≥ 0. Con la elección anterior, Yₜ = Y₀ exp(σWₜ⁽¹⁾ − Wₜ⁽²⁾ − vt/2). Los incrementos gaussianos son independientes del pasado y la corrección −vt/2 hace que su factor exponencial tenga esperanza uno. Por tanto E[Yₜ | ℱₛ] = Yₛ para s ≤ t. Si v = 0, el cociente es constante.')
note(b0,'No basta con afirmar «no hay deriva». La integrabilidad de la exponencial gaussiana, con coeficientes constantes en horizonte finito, justifica que no sea únicamente una martingala local.')

b0=page('Resolver la integral de Itô','Ejercicio 2 · Apartado b')
p(b0,'Queremos expresar la integral usando únicamente el tiempo t y el browniano Wₜ. Seguimos la pista del enunciado: aplicar Itô al cubo del browniano.')
eq(b0,r'I_t=\int_0^t(W_s^2-s)\,dW_s','Integral solicitada; s es la variable de integración temporal.')
eq(b0,r'd(W_t^3)=3W_t^2\,dW_t+3W_t\,dt','Fórmula de Itô para el cubo de un browniano estándar.')
p(b0,'Como W₀ = 0, al integrar y dividir entre tres obtenemos:')
eq(b0,r'\int_0^t W_s^2\,dW_s=\frac{W_t^3}{3}-\int_0^t W_s\,ds','Primera parte de la integral, obtenida a partir de W al cubo.')
p(b0,'Para el segundo término usamos la regla del producto en tWₜ. El tiempo es determinista y no aporta covariación:')
eq(b0,r'd(tW_t)=W_t\,dt+t\,dW_t','Regla del producto para tiempo por browniano.')
eq(b0,r'\int_0^t s\,dW_s=tW_t-\int_0^t W_s\,ds','Segunda parte de la integral.')
p(b0,'Al restar, las dos integrales ordinarias de W se cancelan y queda:')
eq(b0,r'\boxed{I_t=\frac{W_t^3}{3}-tW_t}','Solución cerrada de la integral solicitada.')
h(b0,'Comprobación directa')
p(b0,'Al diferenciar Wₜ³/3 − tWₜ, los términos Wₜ dt se cancelan y queda (Wₜ² − t)dWₜ, exactamente el integrando original. Además, el valor inicial es cero.')
p(b0,'Como control adicional, E[Iₜ] = 0. Por la isometría de Itô, Var(Iₜ) = ∫₀ᵗ E[(Wₛ² − s)²] ds. Usando E[Wₛ²] = s y E[Wₛ⁴] = 3s², el integrando es 2s² y la varianza resulta 2t³/3.')

b0=page('Una martingala cuadrática','Ejercicio 2 · Apartado c.1 y c.2')
p(b0,'Elevar una martingala al cuadrado introduce una deriva positiva. El término t³/3 del enunciado compensa exactamente esa deriva. Para verlo, definimos:')
eq(b0,r'A_t=\int_0^t s\,dW_s,\qquad q(t)=\int_0^t s^2\,ds=\frac{t^3}{3}','A es la integral gaussiana y q su varianza acumulada.')
eq(b0,r'M_t=A_t^2-q(t)','Proceso solicitado, escrito como cuadrado menos varianza.')
h(b0,'Dos vías de demostración')
p(b0,'Primera vía: Itô da d(Aₜ²) = 2tAₜ dWₜ + t² dt. Restar dq(t) = t² dt elimina la deriva, de modo que dMₜ = 2tAₜ dWₜ. Su integrando es cuadrado integrable: E[∫₀ᵀ4t²Aₜ² dt] = 2T⁶/9 < ∞ para todo horizonte finito T.')
p(b0,'Segunda vía, que desarrollamos: calcular la esperanza condicional. Para 0 ≤ s ≤ t, escribimos Aₜ = Aₛ + ΔA. Como el integrando es determinista, ΔA es independiente de la información ℱₛ disponible hasta s, tiene media cero y varianza q(t) − q(s).')
eq(b0,r'E[A_t^2\mid\mathcal{F}_s]=A_s^2+q(t)-q(s)','Esperanza del cuadrado futuro condicionada al pasado.')
p(b0,'El término cruzado 2AₛΔA desaparece al tomar esperanza y E[(ΔA)²] es la varianza del incremento. Por tanto:')
eq(b0,r'E[M_t\mid\mathcal{F}_s]=A_s^2-q(s)=M_s','La esperanza condicional verifica la propiedad de martingala.')
p(b0,'La integrabilidad queda garantizada por E[|Mₜ|] ≤ E[Aₜ²] + q(t) = 2q(t) < ∞; además el proceso es adaptado a la información browniana.')
h(b0,'Distribución y momentos')
p(b0,'Aₜ es normal de media cero y varianza q(t). Para t > 0, Aₜ/√q(t) es normal estándar y su cuadrado tiene distribución chi-cuadrado con un grado de libertad, denotada χ₁².')
eq(b0,r'M_t\ \sim\ \frac{t^3}{3}(\chi_1^2-1)','Ley de M: chi-cuadrado escalada y desplazada.')
eq(b0,r'E[M_t]=0,\qquad \operatorname{Var}(M_t)=\frac{2t^6}{9}','Media y varianza del proceso.')
p(b0,'Se usan E[χ₁²] = 1 y Var(χ₁²) = 2. El soporte es [−t³/3, ∞): no es una normal, sino una distribución asimétrica. En t = 0, M₀ = 0 con probabilidad uno.')

b0=page('Comprobación empírica en Excel','Ejercicio 2 · Apartado c.3')
p(b0,'Una trayectoria muestra cómo evoluciona el proceso; muchas réplicas en t = 1 permiten contrastar su distribución. Son objetivos distintos y el libro contiene ambos.')
h(b0,'Trayectoria en 100 pasos')
p(b0,'Se toman tⱼ = j/100, para j = 0,…,100, y normales estándar independientes Zⱼ. En cada paso se actualiza:')
eq(b0,r'\Delta A_j=\sqrt{\frac{t_j^3-t_{j-1}^3}{3}}\,Z_j','Incremento exacto de A entre dos fechas consecutivas.')
eq(b0,r'A_{t_j}=A_{t_{j-1}}+\Delta A_j,\qquad M_{t_j}=A_{t_j}^2-\frac{t_j^3}{3}','Actualización de la trayectoria de A y de M.')
p(b0,'Se parte de A₀ = M₀ = 0. La varianza de cada incremento se integra exactamente; por tanto no hay error de Euler en las fechas de la malla. La hoja Trayectoria conserva los Z y calcula los pasos mediante fórmulas.')
h(b0,'5.000 réplicas terminales')
p(b0,'En t = 1 basta simular M₁ = (Z² − 1)/3. La hoja Replica_t1 contiene 5.000 normales y las fórmulas; Resumen compara los momentos y el soporte. Los Z están fijados para que recalcular el Excel no cambie la muestra.')
table(b0,['Magnitud','Teoría','Excel'],[
 ['Media','0',num(mean,6)],['Varianza','2/9 = 0,222222',num(var,6)],
 ['Error estándar de la media','√(varianza muestral / 5.000)',num(se,6)],
 ['Mínimo','No menor que −1/3',num(M.min(),9)]])
p(b0,f'El intervalo aproximado del 95 % para la media es media ± 1,96 × error estándar: [{num(mean-1.96*se,6)}; {num(mean+1.96*se,6)}]. Contiene cero; la discrepancia es compatible con error de simulación.')
fig(b0,'cdf.png','Distribución acumulada teórica y empírica en t = 1. Se usan las mismas 5.000 observaciones del Excel, no una muestra adicional.')
p(b0,'La simulación respalda la ley y los momentos; la demostración de martingala es la del apartado anterior. Una media empírica próxima a cero no sustituye a la propiedad de esperanza condicional.')

b0=page('Del modelo a la fórmula de la call','Ejercicio 3 · Apartados 1 y 2')
p(b0,'El enunciado fija S₀ = 100, tipo continuo r = 0,01 y volatilidad determinista σ(t). Interpretamos la dinámica bajo la medida de valoración neutral al riesgo, sin dividendos; los importes se expresan en unidades monetarias.')
eq(b0,r'V(t)=\int_0^t\sigma^2(s)\,ds','V es la varianza acumulada del logaritmo del precio.')
p(b0,'Definimos Lₜ = rt − V(t)/2 + ∫₀ᵗσ(s)dWₛ. Su diferencial es (r − σ²(t)/2)dt + σ(t)dWₜ y su variación cuadrática es σ²(t)dt. Al aplicar Itô a S₀ exp(Lₜ), la corrección del exponencial cancela −σ²(t)/2:')
eq(b0,r'dS_t=S_t\left(dL_t+\frac12d\langle L\rangle_t\right)=rS_t\,dt+\sigma(t)S_t\,dW_t','Verificación de la dinámica del precio propuesto.')
p(b0,'También se cumple la condición inicial. Queda verificada la solución del enunciado:')
eq(b0,r'S_t=S_0\exp\left(rt-\frac12V(t)+\int_0^t\sigma(s)\,dW_s\right)','Solución lognormal con volatilidad determinista.')
h(b0,'La call europea')
p(b0,'La integral gaussiana hasta T tiene media cero y varianza V(T). Así, log S_T tiene media log S₀ + rT − V(T)/2 en el vencimiento T. Para strike K y V(T) > 0 definimos:')
eq(b0,r'd_1=\frac{\log(S_0/K)+rT+V(T)/2}{\sqrt{V(T)}},\qquad d_2=d_1-\sqrt{V(T)}','Umbrales normales de la fórmula de valoración.')
p(b0,'La probabilidad de acabar por encima de K es Φ(d₂). La esperanza truncada del activo es S₀ exp(rT) Φ(d₁), que se obtiene completando el cuadrado en la densidad normal. Descontar ambas partes del payoff produce:')
eq(b0,r'\boxed{C=S_0\Phi(d_1)-Ke^{-rT}\Phi(d_2)}','Precio de la call europea con volatilidad determinista.')
note(b0,'La fórmula es la de Black-Scholes con σ√T sustituido por √V(T). Si σ es constante, V(T) = σ²T y se recupera el caso habitual. Si V(T) = 0, el precio es max(S₀ − K exp(−rT), 0).')

b0=page('Calibrar la volatilidad cuadrática','Ejercicio 3 · Apartado 3')
p(b0,'Los tres precios de call, todos con K = 100, identifican tres varianzas acumuladas. Primero invertimos la fórmula europea en cada vencimiento; después ajustamos los tres coeficientes de σ(t) = at² + bt + c.')
table(b0,['T (años)','Call de mercado','V(T) implícita','Call reconstruida'],[[num(t,2),num(pr,2),num(v,10),num(cf,8)]for t,pr,v,cf in zip(times,prices,Vs,fit)])
p(b0,'Al elevar el polinomio al cuadrado e integrar término a término se obtiene la relación dada en el enunciado:')
eq(b0,r'V(t)=\frac{a^2t^5}{5}+\frac{abt^4}{2}+\frac{(b^2+2ac)t^3}{3}+bct^2+c^2t','Varianza acumulada para la volatilidad cuadrática.')
p(b0,'Se resuelve el sistema V(1) = 0,0521320003, V(1,25) = 0,0899877821 y V(2) = 0,3292216020, usando internamente las cifras sin redondear. La raíz elegida es:')
eq(b0,rf'a={a:.12f},\quad b={b:.12f},\quad c={c:.12f}','Coeficientes calibrados con el tiempo expresado en años.')
p(b0,'La volatilidad se expresa en años⁻¹ᐟ²; por ello las unidades de a, b y c son años⁻⁵ᐟ², años⁻³ᐟ² y años⁻¹ᐟ². Los dígitos extra permiten reproducir el ajuste, pero no implican esa precisión económica en precios de mercado dados con dos decimales.')
fig(b0,'vol.png','Volatilidad instantánea de la solución seleccionada. No debe confundirse con la volatilidad efectiva √[V(T)/T] de una call europea.')
note(b0,'Criterio de modelización: se exige σ(t) ≥ 0 en [0,2]. La búsqueda multiorigen del notebook encuentra cuatro soluciones algebraicas; solo una de las encontradas cumple esta restricción. En la elegida a, b y c son positivos, lo que verifica la positividad en todo el intervalo, no solo en una malla.')
p(b0,'El ajuste europeo no determina por sí solo toda la trayectoria de la volatilidad. La no negatividad es una convención económica explícita, no una condición escrita en el enunciado. Cambiar globalmente el signo de σ no cambia la ley de S; elegir otra forma de σ² entre vencimientos sí puede afectar a la asiática.')

b0=page('Valorar la opción asiática','Ejercicio 3 · Apartado 4')
p(b0,'El payoff depende de la media aritmética de S en 1,15; 1,30; 1,60 y 1,70 años, pero se paga en T = 2. Por eso se simulan las cuatro fijaciones y se descuenta hasta dos años, no hasta la última fijación.')
eq(b0,r'C_A=e^{-0.01\cdot2}E\left[\max\left(\frac{S_{1.15}+S_{1.30}+S_{1.60}+S_{1.70}}{4}-100,0\right)\right]','Valor actual de la call sobre el promedio aritmético de las cuatro fijaciones.')
h(b0,'Simulación exacta en las fechas necesarias')
p(b0,'Tomamos t₀ = 0 y las cuatro fechas tᵢ anteriores. Para cada trayectoria generamos cuatro normales estándar independientes Zᵢ y acumulamos los incrementos de varianza:')
eq(b0,r'B_i=\sum_{j=1}^{i}\sqrt{V(t_j)-V(t_{j-1})}\,Z_j','Integral gaussiana acumulada hasta cada fijación.')
eq(b0,r'S_{t_i}=100\exp\left(0.01t_i-\frac12V(t_i)+B_i\right)','Precio simulado en cada una de las cuatro fechas.')
p(b0,'Bᵢ representa la integral de σ dW, no un nuevo browniano estándar. Compartir los incrementos previos garantiza Cov(Bᵢ,Bⱼ) = V(min(tᵢ,tⱼ)). Simular cada fijación de manera independiente daría un precio incorrecto.')
h(b0,'Estimación y precisión')
p(b0,'La implementación usa puntos Sobol aleatorizados para cubrir el espacio de simulación, parejas antitéticas Z y −Z, y una call sobre la media geométrica como control de varianza. Esta última tiene precio analítico porque el logaritmo de la media geométrica es normal.')
p(b0,'Se realizan 32 aleatorizaciones independientes, cada una con 65.536 parejas (131.072 trayectorias). El precio es la media de las 32 estimaciones; su error estándar se calcula entre réplicas y el intervalo usa el cuantil t de Student con 31 grados de libertad.')
table(b0,['Estimación del notebook','Valor'],[
 ['Precio aritmético','13,621805'],['Error estándar entre réplicas','0,00007890'],
 ['Intervalo aproximado del 95 %','[13,621644; 13,621966]'],
 ['Contraste pseudoaleatorio independiente','13,622091 (error estándar: 0,001054)']])
note(b0,'Resultado para entregar: 13,62 unidades monetarias. El estrecho intervalo mide error de simulación condicionado a la curva calibrada; no mide riesgo de modelo ni incertidumbre de las cotizaciones.')
p(b0,'Como control, el precio geométrico es 12,796771, menor que el aritmético, tal como exige la desigualdad entre medias. Una curva alternativa que cruza cero repricia las calls, pero da aproximadamente 13,740158 para la asiática: el criterio de selección de curva importa.')

b0=page('Detalle del control y comprobaciones','Complemento de reproducibilidad')
h(b0,'Control geométrico: por qué no cambia el precio esperado')
p(b0,'Sean H y G los payoffs descontados aritmético y geométrico de una pareja antitética. Se estima H − β(G − C_G), donde C_G es el precio geométrico exacto. Como E[G] = C_G, la corrección tiene esperanza cero.')
p(b0,'Para reconstruir C_G, el logaritmo del promedio geométrico tiene media m y varianza v_G dadas por:')
eq(b0,r'm=\frac14\sum_{i=1}^{4}\left(\log100+0.01t_i-\frac12V(t_i)\right)','Media del logaritmo del promedio geométrico.')
eq(b0,r'v_G=\frac1{16}\sum_{i=1}^{4}\sum_{j=1}^{4}V(\min(t_i,t_j))','Varianza del logaritmo del promedio geométrico.')
eq(b0,r'd_G=\frac{m-\log100}{\sqrt{v_G}}','Umbral normal del promedio geométrico.')
eq(b0,r'C_G=e^{-0.02}\left[e^{m+v_G/2}\Phi(d_G+\sqrt{v_G})-100\Phi(d_G)\right]','Precio exacto de la call geométrica utilizada como control.')
p(b0,'El notebook obtiene C_G = 12,79677073. Fija β = 1,04122446 mediante una muestra piloto independiente de 200.000 parejas, con β = Cov(H,G)/Var(G). Así, los datos que ajustan el control no son los utilizados para producir el precio final.')
h(b0,'Qué se ha comprobado en esta edición')
p(b0,'Se ha releído el enunciado, recalculado Pⁿ para n = 1,…,25 desde el Excel original y reconstruido la cohorte. Se han contrastado las identidades de Itô y los momentos teóricos. Las 5.000 réplicas del Excel se han verificado contra M₁ = (Z² − 1)/3 y se ha repetido la calibración de las tres calls.')
p(b0,'Las cifras de las simulaciones de cadenas y de la asiática proceden de los notebooks ejecutados del repositorio. Esta edición reorganiza su presentación: no atribuye a una nueva ejecución las simulaciones previamente guardadas.')
h(b0,'Límites que condicionan la interpretación')
p(b0,'La matriz se mantiene constante durante 25 años y Caa-C usa la fila elegida con autorización docente. La volatilidad cuadrática y su no negatividad son elecciones del modelo. Las cifras están redondeadas solo para presentación; las tablas pueden sumar unas milésimas de punto porcentual por encima o por debajo de 100 %.')
p(b0,'El Excel del apartado 2.c se entrega como archivo complementario existente. Su simulación es reproducible con los Z guardados; no es necesario generar nuevos números al abrirlo. En Resumen, la columna Teoría contiene la cota −1/3; la columna Hoja de esa fila repite el mínimo empírico, que no debe interpretarse como una nueva cota teórica.')

b0=page('Probabilidades acumuladas de default','Anexo A · Apartado 1.a')
p(b0,'Fᵢ(n) = (Pⁿ)ᵢD. Cada celda es la probabilidad de haber hecho default como máximo en el año de su fila, condicionada al rating inicial de su columna. Valores en %.')
table(b0,['Año']+RATINGS,[[str(n)]+[num(x*100,4) for x in F[n-1]] for n in range(1,26)])
p(b0,'0,0000 significa que el valor redondea a cero, no necesariamente que sea exactamente nulo. Las probabilidades aumentan con el horizonte porque Default es absorbente.')

b0=page('Probabilidades de primer default','Anexo B · Apartado 1.a')
p(b0,'fᵢ(n) = Fᵢ(n) − Fᵢ(n−1). Cada celda es la probabilidad de hacer default exactamente en ese año, respecto al grupo inicialmente en ese rating. Valores en %.')
table(b0,['Año']+RATINGS,[[str(n)]+[num(x*100,4) for x in f[n-1]] for n in range(1,26)])
p(b0,'La suma de cada columna coincide con el acumulado a 25 años del anexo A, antes de redondear. No es la probabilidad condicionada a haber sobrevivido hasta el año anterior.')

b0=page('Defaults esperados de la cohorte','Anexo C · Apartado 1.b')
p(b0,'La proporción anual g(n) se pondera con la composición inicial de las 5.002 compañías. El acumulado es la suma de las proporciones anuales hasta ese año.')
table(b0,['Año','Default en ese año (%)','Compañías esperadas','Acumulado (%)'],[[str(n),num(g[n-1]*100,4),num(5002*g[n-1],2),num(g[:n].sum()*100,4)]for n in range(1,26)])
note(b0,'Conclusión: los tres ejercicios distinguen resultado exacto y estimación. En ratings se encadenan transiciones; en Itô se identifica y compensa la deriva; en opciones se calibra la varianza acumulada y se simulan fijaciones correlacionadas con descuento al pago.')

CSS='''
:root{color-scheme:light;--ink:#20332f;--muted:#596b65;--accent:#175b54;--line:#d9e0dc;--paper:#fff;--soft:#f3f6f3}
*{box-sizing:border-box}body{margin:0;background:#f6f7f4;color:var(--ink);font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif;font-size:17px;line-height:1.62;-webkit-font-smoothing:antialiased}
main{max-width:880px;margin:32px auto 70px;background:var(--paper);padding:60px 66px;box-shadow:0 1px 20px #20332f08}
.kicker{font-size:12px;letter-spacing:.13em;text-transform:uppercase;color:var(--accent);font-weight:650;margin:0 0 10px}h1{font-family:Georgia,serif;font-size:48px;font-weight:400;line-height:1.1;letter-spacing:-.035em;margin:0 0 30px}h2{font-family:Georgia,serif;font-weight:400;font-size:32px;line-height:1.22;letter-spacing:-.02em;margin:0 0 24px}h3{font-size:18px;margin:28px 0 10px}p{margin:0 0 17px}section{padding-top:44px;margin-top:44px;border-top:1px solid var(--line)}section:first-of-type{padding:0;margin:0;border:0}.note{border-left:3px solid var(--accent);background:var(--soft);padding:17px 20px;margin:23px 0}.formula{overflow-x:auto;max-width:100%;padding:18px 4px;margin:8px 0 20px;-webkit-overflow-scrolling:touch}math{font-size:1.05em}figure{margin:24px 0}img{display:block;max-width:100%;height:auto}figcaption{font-size:13px;line-height:1.5;color:var(--muted);margin-top:10px}.table-scroll{overflow-x:auto;max-width:100%;-webkit-overflow-scrolling:touch;margin:12px 0 22px}table{border-collapse:collapse;width:100%;font-size:14px;line-height:1.45;font-variant-numeric:tabular-nums}td,th{padding:11px 10px;text-align:left;border-bottom:1px solid var(--line);vertical-align:top}thead th{font-size:12px;color:var(--accent);background:var(--soft);font-weight:650}.wide{min-width:690px}.wide td,.wide th{padding:9px 7px;white-space:nowrap}.swipe-hint{font-size:12px;color:var(--muted);margin:0}details{border-bottom:1px solid var(--line);margin-bottom:30px}summary{min-height:44px;padding:12px 0;cursor:pointer;color:var(--accent)}nav a{display:block;padding:10px 0;min-height:44px;color:var(--accent);text-decoration:none}footer{margin-top:40px;color:var(--muted);font-size:13px}a:focus-visible,summary:focus-visible{outline:2px solid var(--accent);outline-offset:4px}
@media(max-width:600px){body{font-size:16.5px;line-height:1.58}main{margin:0;padding:32px max(19px,env(safe-area-inset-right)) 40px max(19px,env(safe-area-inset-left));width:100%;box-shadow:none}h1{font-size:39px}h2{font-size:28px}section{margin-top:34px;padding-top:34px}.note{padding:14px 15px}td,th{padding:9px 7px}table{font-size:13px}}
@media(prefers-color-scheme:dark){:root{color-scheme:dark;--ink:#e8eeea;--muted:#b4c4bd;--accent:#9ac9ba;--line:#3c5149;--paper:#192620;--soft:#24362e}body{background:#142019}img{background:white}}
@media print{body{background:white;font-size:11pt}main{max-width:none;margin:0;padding:0;box-shadow:none}nav,.swipe-hint{display:none}section{break-before:page;border:0}section:first-of-type{break-before:auto}.formula,.table-scroll{overflow:visible}.wide{min-width:0}h2,h3{break-after:avoid}figure,table,.note{break-inside:avoid}}
'''
for fname,name in [('DejaVuSans.ttf','Body'),('DejaVuSans-Bold.ttf','Bold'),('DejaVuSerif.ttf','Serif')]:
    pdfmetrics.registerFont(TTFont(name,'/usr/share/fonts/truetype/dejavu/'+fname))
pdfmetrics.registerFontFamily('Body',normal='Body',bold='Bold',italic='Body',boldItalic='Bold')
styles={
 'p':ParagraphStyle('p',fontName='Body',fontSize=9.2,leading=13.7,spaceAfter=9,textColor=HexColor('#20332f')),
 'h':ParagraphStyle('h',fontName='Bold',fontSize=11,leading=15,spaceBefore=11,spaceAfter=7,keepWithNext=True,textColor=HexColor('#175b54')),
 'title':ParagraphStyle('title',fontName='Serif',fontSize=24,leading=29,spaceAfter=19,textColor=HexColor('#20332f')),
 'kicker':ParagraphStyle('kicker',fontName='Bold',fontSize=8,leading=11,spaceAfter=10,textColor=HexColor('#175b54')),
 'caption':ParagraphStyle('caption',fontName='Body',fontSize=7.8,leading=11,spaceAfter=10,textColor=HexColor('#596b65')),
 'td':ParagraphStyle('td',fontName='Body',fontSize=8,leading=11,textColor=HexColor('#20332f')),
 'smalltd':ParagraphStyle('smalltd',fontName='Body',fontSize=7.1,leading=10,textColor=HexColor('#20332f')),
 'th':ParagraphStyle('th',fontName='Bold',fontSize=7.1,leading=10,textColor=HexColor('#175b54')),
}
WIDTH=A4[0]-108
def para(s,style='p'):return Paragraph(escape(s),styles[style])
story=[]
html=['<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover"><title>MEFC 2026 · Procesos estocásticos</title><style>'+CSS+'</style></head><body><main>']
nav='<nav aria-label="Índice"><details><summary>Contenido de la memoria</summary>'+''.join(f'<a href="#seccion-{i}">{escape(k)} · {escape(t)}</a>'for i,(t,k,_)in enumerate(pages))+'</details></nav>'
equation_count=0
for i,(title,kicker,blocks) in enumerate(pages):
    if i:story.append(PageBreak())
    story.extend([para(kicker.upper(),'kicker'),para(title,'title')])
    html.append(f'<section id="seccion-{i}"><header><p class="kicker">{escape(kicker)}</p><h{1 if i==0 else 2}>{escape(title)}</h{1 if i==0 else 2}></header>')
    if i==0:html.append(nav)
    for item in blocks:
        kind=item[0]
        if kind in ('p','h'):
            tag='p' if kind=='p' else 'h3'
            html.append(f'<{tag}>{escape(item[1])}</{tag}>')
            story.append(para(item[1],kind))
        elif kind=='note':
            html.append('<aside class="note">'+escape(item[1])+'</aside>')
            box=Table([[para(item[1])]],colWidths=[WIDTH])
            box.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),HexColor('#f3f6f3')),('LINEBEFORE',(0,0),(0,-1),2,HexColor('#175b54')),('LEFTPADDING',(0,0),(-1,-1),12),('TOPPADDING',(0,0),(-1,-1),9),('BOTTOMPADDING',(0,0),(-1,-1),3)]))
            story.extend([box,Spacer(1,10)])
        elif kind=='eq':
            tex,label=item[1:]
            mml=latex2mathml.converter.convert(tex,display='block').replace('<math ',f'<math aria-label="{escape(label,quote=True)}" ',1)
            html.append('<div class="formula" tabindex="0">'+mml+'</div>')
            # mathtext does not implement boxed; color and spacing provide emphasis.
            pdftex=tex.replace(r'\boxed{',r'\mathrm{').replace(r'\operatorname{Var}',r'\mathrm{Var}').replace(r'\frac12',r'\frac{1}{2}').replace(r'\frac14',r'\frac{1}{4}').replace(r'\frac1{16}',r'\frac{1}{16}')
            buf=BytesIO();math_to_image('$'+pdftex+'$',buf,prop=FontProperties(size=13),dpi=270,format='png',color='#20332f')
            buf.seek(0)
            im=PILImage.open(buf);iw,ih=im.size
            target=min(WIDTH,iw*72/270)
            equation=Image(buf,width=target,height=ih*target/iw)
            story.extend([Spacer(1,4),equation,Spacer(1,13)])
            equation_count+=1
        elif kind=='table':
            headers,rows=item[1:];wide=len(headers)>4
            html.append('<p class="swipe-hint">Desliza para ver más →</p><div class="table-scroll" tabindex="0"><table'+(' class="wide"'if wide else '')+'><thead><tr>'+''.join('<th scope="col">'+escape(x)+'</th>'for x in headers)+'</tr></thead><tbody>')
            for row in rows:html.append('<tr>'+''.join(('<th scope="row">'if j==0 else '<td>')+escape(str(x))+('</th>'if j==0 else '</td>')for j,x in enumerate(row))+'</tr>')
            html.append('</tbody></table></div>')
            data=[[para(str(x),'th') for x in headers]]+[[para(str(x),'smalltd'if wide else 'td') for x in row]for row in rows]
            widths=[WIDTH/len(headers)]*len(headers)
            if len(headers)==2:widths=[WIDTH*.51,WIDTH*.49]
            if len(headers)==3:widths=[WIDTH*.30,WIDTH*.34,WIDTH*.36]
            tb=Table(data,colWidths=widths,repeatRows=1,hAlign='LEFT')
            tb.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),HexColor('#f3f6f3')),('LINEBELOW',(0,0),(-1,-1),.35,HexColor('#d9e0dc')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]))
            story.extend([tb,Spacer(1,12)])
        elif kind=='fig':
            path=TMP/item[1];caption=item[2]
            encoded=base64.b64encode(path.read_bytes()).decode()
            html.append(f'<figure><img src="data:image/png;base64,{encoded}" alt="{escape(caption,quote=True)}"><figcaption>{escape(caption)}</figcaption></figure>')
            im=PILImage.open(path);iw,ih=im.size
            target=min(WIDTH,{'cdf.png':350,'vol.png':380}.get(item[1],430))
            story.append(KeepTogether([Image(str(path),width=target,height=ih*target/iw),Spacer(1,6),para(caption,'caption')]))
    html.append('</section>')
html.append('<footer>MEFC 2026 · Procesos estocásticos · Fin de la memoria</footer></main></body></html>')
(OUT/'main.html').write_text('\n'.join(html),encoding='utf-8')
def footer(canvas,doc):
    canvas.saveState();canvas.setFont('Body',7.5);canvas.setFillColor(HexColor('#596b65'))
    canvas.drawString(54,28,'MEFC 2026 · Procesos estocásticos')
    canvas.drawRightString(A4[0]-54,28,str(doc.page));canvas.restoreState()
doc=SimpleDocTemplate(str(OUT/'procesos_estocasticos.pdf'),pagesize=A4,leftMargin=54,rightMargin=54,topMargin=43,bottomMargin=48,title='MEFC 2026 · Procesos estocásticos',author='Grupo MEFC 2026')
doc.build(story,onFirstPage=footer,onLaterPages=footer)
report={'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'matrix_row_error':float(max(abs(P.sum(1)-1))),'cohort_first_year':float(g[0]),'cohort_25_years':float(g.sum()),'excel_mean':float(mean),'excel_variance':float(var),'calibration':abc.tolist(),'repricing_error':float(max(abs(fit-prices))),'sections':len(pages),'equations':equation_count}
(TMP/'checks.json').write_text(json.dumps(report,indent=2))
print(json.dumps(report,indent=2))
print(OUT/'main.html');print(OUT/'procesos_estocasticos.pdf')
